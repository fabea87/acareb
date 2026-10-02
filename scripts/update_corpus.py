#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AcadRebuttal - 具高度扩展性与前瞻性的自适应学术语料发现与编译流水线 (Next-Gen Auto-Discovery)
支持未来任意新增论文工作区、多轮返修文件夹、不同 Word/PDF 排版风格的自动嗅探与智能抽取。
"""

import os
import sys
import io
import re
import json
import subprocess
import docx
import pypdf

# Force UTF-8 on Windows
if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WORKSPACE_DIR = os.path.dirname(PROJECT_DIR)
OUTPUT_JSON = os.path.join(PROJECT_DIR, "src", "data", "corpus.json")

# Ignored directory names (build artifacts, system folders, git, etc.)
IGNORED_DIRS = {
    'acad-rebuttal-web', 'node_modules', '.git', '.vscode', '.idea', 'dist',
    'review_corpus_system', '__pycache__', 'build', '.obsidian'
}

# Journal Knowledge Base for automatic normalization & tier enrichment
JOURNAL_KNOWLEDGE_BASE = {
    "system": ("System", "SSCI Q1 · Elsevier"),
    "call": ("Computer Assisted Language Learning (CALL)", "SSCI Q1 · Taylor & Francis"),
    "llt": ("Language Learning & Technology (LLT)", "SSCI Q1 · Open Access"),
    "eait": ("Education and Information Technologies (EAIT)", "SSCI Q1 · Springer"),
    "ile": ("Interactive Learning Environments (ILE)", "SSCI Q1 · Taylor & Francis"),
    "hssc": ("Humanities and Social Sciences Communications (HSSC)", "SSCI Q1 · Nature Portfolio"),
    "kybernetes": ("Kybernetes", "SSCI / SCI · Emerald"),
    "perspectives": ("Perspectives: Studies in Translation Theory and Practice", "A&HCI / SSCI Q1 · Routledge"),
    "plos one": ("PLOS ONE", "SCI / SSCI Q1 · PLOS"),
    "plos": ("PLOS ONE", "SCI / SSCI Q1 · PLOS"),
    "injal": ("International Journal of Applied Linguistics (InJAL)", "SSCI · Wiley"),
    "fip": ("Frontiers in Psychology (Fip)", "SSCI Q1 · Frontiers"),
    "frontiers": ("Frontiers in Psychology (Fip)", "SSCI Q1 · Frontiers"),
    "caeh": ("Assessment & Evaluation in Higher Education (CAEH)", "SSCI Q1 · Taylor & Francis"),
    "aehe": ("Assessment & Evaluation in Higher Education (CAEH)", "SSCI Q1 · Taylor & Francis"),
    "ijarped": ("International Journal of Academic Research in Progressive Education and Development", "Peer-Reviewed")
}

# 8 Core Applied Linguistics & Educational Technology Issue Categories
CATEGORIES = [
    ("Theory & Framing", ["theor", "framework", "concept", "construct", "underpin", "grounding", "rationale", "conceptual", "paradigm", "lens"]),
    ("Methodology & Design", ["method", "design", "sample", "participant", "intervention", "procedure", "group", "control", "quasi-experiment", "setting", "cohort", "protocol", "task", "fidelity"]),
    ("Measurement & Rubrics", ["instrument", "measure", "questionnaire", "survey", "scale", "rubric", "validity", "reliability", "proficiency", "pre-test", "post-test", "score", "dimension", "tem-4"]),
    ("Literature & Scoping", ["literature", "cite", "citation", "reference", "previous studies", "prior work", "scope", "relevant", "suggested paper", "recent research", "scholar"]),
    ("Data Analysis & Stats", ["statistic", "analysis", "data", "p-value", "effect size", "anova", "regression", "model", "sem", "ena", "coding", "qualitative", "quantitative", "variable", "correlation", "normality", "kappa"]),
    ("Discussion & Argumentation", ["discuss", "discussion", "implication", "limitation", "mechanism", "explain", "interpret", "critical", "claim", "argumentation", "generalize", "future work"]),
    ("GenAI & Tech Implementation", ["chatgpt", "prompt", "ai", "genai", "automated", "algorithm", "teqsa", "policy", "llm", "technology-enhanced", "digital", "bot"]),
    ("Minor Revision & Style", ["typo", "wording", "style", "tone", "table", "figure", "format", "title", "grammar", "line", "page", "redundancy", "clarification", "spelling", "syntax"])
]

STRATEGIES = [
    ("Polite Defense & Boundary", ["while we", "outside the scope", "beyond the scope", "focus of this study", "deliberately", "purposefully", "not the primary aim", "however, we", "it should be noted that some", "contrary to"]),
    ("Clarify & Reframe", ["clarif", "distinguish", "delineate", "we clarify", "misunderstand", "convey", "we emphasize", "definition", "term use", "operationaliz"]),
    ("Methodological Fortification", ["power", "robustness", "validity", "reliability", "pre-assessment", "triangulat", "detailed description", "justif", "rigor", "pilot", "baseline"]),
    ("Empirical Supplement", ["statistical", "table", "figure", "additional data", "coding", "transcript", "anova", "regression", "evidence", "appendix"]),
    ("Direct Correction", ["fixed", "corrected", "rectified", "typo", "applied", "changed to", "modified", "spelling", "grammar"]),
    ("Concede & Expand", ["agree", "have added", "have expanded", "have incorporated", "expanded", "supplemented", "included as suggested", "we added", "rearranged"])
]

def infer_journal_and_topic(folder_name):
    """
    Intelligently infer normalized journal name, tier, and paper topic from folder name.
    Supports formats like:
      - 'System - GenAI助力同伴反馈的效果'
      - '2026.11 CALL - 自动写作反馈与学习者自主性'
      - 'Perspectives - 同传工作压力'
    """
    clean_name = re.sub(r'^[0-9]{4}[.\-_][0-9]{1,2}\s*', '', folder_name).strip()
    
    parts = re.split(r'[-–—:]\s*', clean_name, maxsplit=1)
    raw_journal = parts[0].strip()
    topic = parts[1].strip() if len(parts) > 1 else clean_name

    # Normalize journal via knowledge base
    raw_lower = raw_journal.lower()
    matched_journal = None
    matched_tier = "Peer-Reviewed Academic Journal"

    for key, (norm_name, tier) in JOURNAL_KNOWLEDGE_BASE.items():
        if key in raw_lower:
            matched_journal = norm_name
            matched_tier = tier
            break

    if not matched_journal:
        matched_journal = raw_journal.upper() if len(raw_journal) <= 6 else raw_journal

    return matched_journal, matched_tier, topic

def infer_revision_round(file_path):
    """
    Infer revision round (Rev #1, Rev #2, Rev #3, etc.) from path or filename.
    """
    norm_path = file_path.replace('\\', '/').lower()
    match = re.search(r'\b(?:rev|round|r)\s*(?:#|no\.?)?\s*([0-9]+)\b', norm_path)
    if match:
        return f"Rev #{match.group(1)}"
    if "tape" in norm_path:
        return "Rev #1 (TAPE)"
    return "Rev #1"

def is_rebuttal_file(filename):
    """
    Determine if a file is a candidate response/rebuttal document.
    """
    lower = filename.lower()
    if not lower.endswith(('.docx', '.pdf')):
        return False
        
    # Ignore cover letters, standalone manuscripts, proof sheets, surveys, references
    if any(k in lower for k in ['cover letter', 'manuscript', 'proof', 'survey', 'appendix', 'fixed_ref', '最终出版', 'camarata']):
        return False

    keywords = ['response', 'reviewer', 'rebuttal', 'reply', 'revision', 'comment', '意见', '回复', '修改说明', '返修']
    return any(k in lower for k in keywords)

def scan_workspace():
    """
    Auto-discover all paper directories and candidate response files.
    """
    discovered = []
    
    for item in sorted(os.listdir(WORKSPACE_DIR)):
        full_dir = os.path.join(WORKSPACE_DIR, item)
        if not os.path.isdir(full_dir) or item in IGNORED_DIRS or item.startswith('.'):
            continue

        journal, tier, topic = infer_journal_and_topic(item)
        target_files = []

        # Recursively search for rebuttal files
        for root, dirs, files in os.walk(full_dir):
            for f in files:
                if is_rebuttal_file(f):
                    full_fpath = os.path.join(root, f)
                    rel_fpath = os.path.relpath(full_fpath, WORKSPACE_DIR)
                    round_name = infer_revision_round(rel_fpath)
                    target_files.append((round_name, full_fpath, rel_fpath))

        # Deduplicate files (prefer docx over pdf if same round has both)
        rounds_seen = {}
        for r_name, full_p, rel_p in target_files:
            if r_name not in rounds_seen:
                rounds_seen[r_name] = (full_p, rel_p)
            else:
                existing_full, _ = rounds_seen[r_name]
                if existing_full.endswith('.pdf') and full_p.endswith('.docx'):
                    rounds_seen[r_name] = (full_p, rel_p)

        final_files = [(r_name, f_full, f_rel) for r_name, (f_full, f_rel) in rounds_seen.items()]

        if final_files:
            discovered.append({
                "folder": item,
                "journal": journal,
                "journal_tier": tier,
                "topic": topic,
                "files": final_files
            })

    return discovered

def classify_text(comment, response):
    full = (comment + " " + response).lower()
    best_cat = "Methodology & Design"
    best_score = -1
    for cat, kws in CATEGORIES:
        score = sum(1 for kw in kws if kw in full)
        if score > best_score and score > 0:
            best_score = score
            best_cat = cat
            
    resp_lower = response.lower()
    best_strat = "Concede & Expand"
    best_strat_score = -1
    for strat, kws in STRATEGIES:
        score = sum(1 for kw in kws if kw in resp_lower)
        if score > best_strat_score and score > 0:
            best_strat_score = score
            best_strat = strat
            
    return best_cat, best_strat

def extract_golden_move(response):
    cleaned = re.sub(r'^(?:author(?:\'s)?\s+)?response(?:\s+to\s+reviewer)?\s*[:：\-]\s*', '', response.strip(), flags=re.I)
    sentences = [s.strip().replace("\n", " ") for s in re.split(r'(?<=[.!?])\s+', cleaned) if s.strip()]
    
    patterns = [
        r'\b(?:in the revis|in the updated|we have added|we have expanded|we have revised|we incorporated|we clarified|to address this|we rearranged|we conducted|we purposely|we have introduced|we carefully checked)\b',
        r'\b(?:we agree with the reviewer that|we agree that|while we acknowledge|it should be noted that|this distinction is made because|we appreciate the reviewer\'s suggestion)\b'
    ]
    for pat in patterns:
        for s in sentences:
            if 30 < len(s) < 280 and re.search(pat, s, re.I):
                return s
                
    for s in sentences:
        if 40 < len(s) < 260 and not s.lower().startswith(('thanks for', 'thank you for')):
            return s
            
    if sentences:
        return sentences[0][:180]
    return cleaned[:160] + "..."

def extract_keywords(text):
    words = re.findall(r'\b[a-zA-Z]{4,}\b', text.lower())
    stopwords = {
        'this', 'that', 'with', 'from', 'have', 'were', 'been', 'which', 'their', 'there',
        'about', 'more', 'also', 'some', 'these', 'would', 'could', 'should', 'than', 'into',
        'study', 'paper', 'manuscript', 'reviewer', 'author', 'comment', 'response', 'thanks',
        'thank', 'point', 'suggest', 'suggestion', 'section', 'details'
    }
    freq = {}
    for w in words:
        if w not in stopwords:
            freq[w] = freq.get(w, 0) + 1
    sorted_words = sorted(freq.items(), key=lambda x: x[1], reverse=True)
    return [w for w, _ in sorted_words[:5]]

def parse_docx_adaptive(fpath):
    try:
        doc = docx.Document(fpath)
    except Exception as e:
        print(f"    [警告] 无法读取 Docx 文件: {fpath} ({e})")
        return []

    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    
    # 1. Parse tables if present (common in CALL, System, etc.)
    table_items = []
    for table in doc.tables:
        if not table.rows:
            continue
        header = [c.text.strip().lower() for c in table.rows[0].cells]
        c_idx, r_idx = -1, -1
        for idx, text in enumerate(header):
            if any(k in text for k in ['comment', 'request', 'suggestion', 'issue', 'item', 'review', '意见', '问题']):
                c_idx = idx
            elif any(k in text for k in ['response', 'reply', 'action', 'author', '回复', '修改说明']):
                r_idx = idx
        if c_idx != -1 and r_idx != -1 and c_idx != r_idx:
            for row in table.rows[1:]:
                cells = [c.text.strip() for c in row.cells]
                if len(cells) > max(c_idx, r_idx):
                    comm = cells[c_idx]
                    resp = cells[r_idx]
                    if comm and resp and len(comm) > 3:
                        table_items.append(("Reviewer", comm, resp))
                        
    # 2. Parse paragraph Q&A stream
    para_items = []
    current_rev = "Reviewer"
    curr_c = []
    curr_r = []
    in_resp = False
    
    rev_pattern = re.compile(r'^(?:(?:comments?\s+(?:from|by|of))|(?:response\s+to\s+(?:the\s+)?comments?\s+(?:from|by|of))|(?:referee|reviewer|editor))\s*(?:#|no\.?)?\s*([0-9a-zA-Z]+|\b(?:one|two|three|editor|associate\s+editor)\b)?', re.I)
    resp_start_pattern = re.compile(r'^(?:author(?:\'s)?\s+)?response(?:\s+to\s+reviewer)?\s*[:：\-]', re.I)
    numbered_pattern = re.compile(r'^(?:(?:comment\s*)?[0-9]{1,2}[.)]|Q[0-9]{1,2}[:.]|\([0-9]{1,2}\))\s+', re.I)
    
    for p in paras:
        if rev_pattern.search(p) and len(p) < 80 and not resp_start_pattern.search(p):
            if curr_c and curr_r:
                para_items.append((current_rev, "\n".join(curr_c), "\n".join(curr_r)))
                curr_c, curr_r = [], []
            current_rev = p
            in_resp = False
            continue
            
        if resp_start_pattern.search(p):
            in_resp = True
            curr_r.append(p)
            continue
            
        if in_resp:
            if numbered_pattern.match(p) or (len(p) < 80 and p.lower().startswith('comment')):
                if curr_c and curr_r:
                    para_items.append((current_rev, "\n".join(curr_c), "\n".join(curr_r)))
                    curr_c, curr_r = [], []
                in_resp = False
                curr_c.append(p)
            else:
                curr_r.append(p)
        else:
            if current_rev == "Reviewer" and any(k in p.lower() for k in ['dear editor', 'we would like to thank', 'greetings']):
                continue
            curr_c.append(p)
            
    if curr_c and curr_r:
        para_items.append((current_rev, "\n".join(curr_c), "\n".join(curr_r)))
        
    return para_items + table_items

def parse_pdf_adaptive(fpath):
    try:
        reader = pypdf.PdfReader(fpath)
    except Exception as e:
        print(f"    [警告] 无法读取 PDF 文件: {fpath} ({e})")
        return []

    lines = []
    for p in reader.pages:
        txt = p.extract_text() or ""
        lines.extend([l.strip() for l in txt.split('\n') if l.strip()])
        
    items = []
    curr_c = []
    curr_r = []
    in_resp = False
    
    for l in lines:
        if any(k in l.lower() for k in ['dear reviewer', 'updated manuscript is appended', 'manuscript title']):
            continue
        if any(l.startswith(k) for k in ['TERM USE', 'INTRODUCTION', 'METHOD', 'RESULTS', 'DISCUSSION']):
            if curr_c and curr_r:
                items.append(('Reviewer 1', '\n'.join(curr_c), '\n'.join(curr_r)))
                curr_c, curr_r = [], []
            in_resp = False
            continue
            
        if re.match(r'^[0-9]+\.\s+', l):
            if curr_c and curr_r:
                items.append(('Reviewer 1', '\n'.join(curr_c), '\n'.join(curr_r)))
                curr_c, curr_r = [], []
            in_resp = False
            curr_c.append(l)
            continue
            
        if any(l.lower().startswith(k) for k in ['thanks for pointing', 'thanks for the comment', 'sorry for missing', 'we truly agree', 'thanks for the suggestion', 'response:']):
            in_resp = True
            curr_r.append(l)
            continue
            
        if in_resp:
            curr_r.append(l)
        else:
            curr_c.append(l)
            
    if curr_c and curr_r:
        items.append(('Reviewer 1', '\n'.join(curr_c), '\n'.join(curr_r)))
    return items

def build_knowledge_graph(records):
    nodes = []
    node_id_map = {}
    links = []

    def add_node(nid, label, ntype, group, size=15, meta=None):
        if nid not in node_id_map:
            node = {
                "id": nid,
                "label": label,
                "type": ntype,
                "group": group,
                "size": size,
                "meta": meta or {}
            }
            nodes.append(node)
            node_id_map[nid] = node
        else:
            node_id_map[nid]["size"] += 1
        return nid

    journal_counts = {}
    for r in records:
        j = r['journal']
        journal_counts[j] = journal_counts.get(j, 0) + 1

    for j, cnt in journal_counts.items():
        short_label = j.split('(')[-1].replace(')', '') if '(' in j else j[:18]
        add_node(f"j_{j}", short_label, "journal", 1, size=24 + cnt, meta={"full_name": j, "count": cnt})

    cat_counts = {}
    for r in records:
        c = r['category']
        cat_counts[c] = cat_counts.get(c, 0) + 1

    for c, cnt in cat_counts.items():
        add_node(f"cat_{c}", c, "category", 2, size=22 + cnt, meta={"count": cnt})

    strat_counts = {}
    for r in records:
        s = r['strategy']
        strat_counts[s] = strat_counts.get(s, 0) + 1

    for s, cnt in strat_counts.items():
        add_node(f"strat_{s}", s, "strategy", 3, size=20 + cnt // 2, meta={"count": cnt})

    kw_freq = {}
    for r in records:
        for kw in r.get('keywords', []):
            kw_freq[kw] = kw_freq.get(kw, 0) + 1
    
    top_kws = sorted(kw_freq.items(), key=lambda x: x[1], reverse=True)[:18]
    for kw, cnt in top_kws:
        add_node(f"kw_{kw}", kw, "keyword", 4, size=14 + cnt, meta={"count": cnt})

    j_cat_links = {}
    for r in records:
        key = (f"j_{r['journal']}", f"cat_{r['category']}")
        j_cat_links[key] = j_cat_links.get(key, 0) + 1

    for (source, target), weight in j_cat_links.items():
        links.append({"source": source, "target": target, "value": weight, "type": "journal_category"})

    cat_strat_links = {}
    for r in records:
        key = (f"cat_{r['category']}", f"strat_{r['strategy']}")
        cat_strat_links[key] = cat_strat_links.get(key, 0) + 1

    for (source, target), weight in cat_strat_links.items():
        links.append({"source": source, "target": target, "value": weight, "type": "category_strategy"})

    for r in records:
        for kw in r.get('keywords', []):
            if f"kw_{kw}" in node_id_map:
                links.append({"source": f"cat_{r['category']}", "target": f"kw_{kw}", "value": 1, "type": "category_keyword"})

    link_dict = {}
    for l in links:
        k = (l['source'], l['target'])
        if k not in link_dict:
            link_dict[k] = l
        else:
            link_dict[k]['value'] += l['value']

    return {
        "nodes": nodes,
        "links": list(link_dict.values()),
        "stats": {
            "total_items": len(records),
            "total_nodes": len(nodes),
            "total_links": len(link_dict)
        }
    }

def get_strategy_vault(records):
    return [
        {
            "category": "Theory & Framing",
            "challenge": "审稿人质疑理论框架单薄、概念界定不清或概念间逻辑链条断裂",
            "core_move": "承认歧义 -> 澄清概念外延 -> 在正文开辟独立小节重申理论透镜 -> 引用权威元理论支撑",
            "template_en": "We appreciate the reviewer's critical eye on the theoretical underpinnings. We agree that the distinction between [Concept A] and [Concept B] required a sharper operational definition. In the revised manuscript, Section 2.1 has been substantially rewritten to explicitly ground our design in [Author, Year]'s framework...",
            "examples_count": sum(1 for r in records if r['category'] == 'Theory & Framing')
        },
        {
            "category": "Literature & Scoping",
            "challenge": "审稿人要求增加其推荐的一批文献，或质疑文献综述遗漏特定领域",
            "core_move": "表达感谢 -> 逐一审读所提文献 -> 纳入直接相关研究并在综述中对话 -> 对偏离焦点文献礼貌划定本研究边界",
            "template_en": "Thanks for the comment. We have carefully checked all of the suggested papers for addition. We were mostly grateful to identify [3-4 papers] directly pertinent to our scope, which have been thoroughly discussed in Section 1.2. It should be noted that while [other suggested papers] offer valuable insights, they primarily focus on [...], which falls outside the targeted scope of our empirical investigation...",
            "examples_count": sum(1 for r in records if r['category'] == 'Literature & Scoping')
        },
        {
            "category": "Methodology & Design",
            "challenge": "审稿人质疑样本量代表性、前测有效性、分组控制或干预程序细节不足",
            "core_move": "肯定方法严谨性关切 -> 补全前测分数检验与准实验基线齐同性 -> 增补详尽干预时间线与任务规范附录",
            "template_en": "We appreciate your focus on methodological rigor. In the revision, we have added more details regarding participant recruitment, language proficiency baseline (e.g., TEM-4 standardized scores), and the exact timeline of the intervention in Section 3.2. As for the sample size, we have elaborated on the ecological validity in classroom-based quasi-experiments...",
            "examples_count": sum(1 for r in records if r['category'] == 'Methodology & Design')
        },
        {
            "category": "Data Analysis & Stats",
            "challenge": "审稿人要求补充稳健性检验、正态性说明、效应量（Effect Size）或质性编码信度",
            "core_move": "直接执行补充计算 -> 在答复信中列出新统计表 -> 说明正文已同步更新并标注位置",
            "template_en": "Thanks for the insightful inquiry. In response, we conducted supplementary robustness checks (including normality diagnostics and Cohen's d effect sizes). The complete ANOVA summary table and inter-coder reliability (Cohen's Kappa = .86) have now been incorporated into Table 3 on page 14...",
            "examples_count": sum(1 for r in records if r['category'] == 'Data Analysis & Stats')
        },
        {
            "category": "GenAI & Tech Implementation",
            "challenge": "审稿人关注 AI Prompt 提示词未公开、高校 AI 伦理政策模糊或技术介入的黑箱问题",
            "core_move": "将完整 Prompt 整理为独立附录 -> 阐释制度伦理政策（如 TEQSA）-> 明确强调以学习者为中心的辅助定位",
            "template_en": "We agree with the reviewer that transparency in prompt design is crucial for reproducibility. In the revision, the full prompt templates utilized for AI feedback generation are now provided in Appendix 1. Furthermore, we expanded the subsection on institutional GenAI policy (echoing TEQSA guidelines) to clarify human-AI collaborative agency...",
            "examples_count": sum(1 for r in records if r['category'] == 'GenAI & Tech Implementation')
        },
        {
            "category": "Discussion & Argumentation",
            "challenge": "审稿人认为讨论部分只复述结果，缺乏批判性机制阐释或局限性不够坦诚",
            "core_move": "将孤立数据上升到认知/社会心理学机制解释 -> 坦诚拓展 Limitation 并化为未来研究建议",
            "template_en": "We appreciate this constructive suggestion. Rather than simply restating the quantitative gains, the revised Discussion section critically interprets why the collaborative condition outperformed the individual setting through the lens of cognitive load and socially shared regulation. Additionally, we explicitly acknowledge the longitudinal constraints in the Limitations...",
            "examples_count": sum(1 for r in records if r['category'] == 'Discussion & Argumentation')
        }
    ]

def get_presubmission_scanner_data():
    return {
        "total_paradigms": 6,
        "paradigms": [
            {
                "id": "quasi_experiment",
                "name": "准实验 / 课堂干预设计 (Quasi-experimental / Intervention)",
                "icon": "flask-conical",
                "tag": "实证干预必测",
                "description": "二语写作与教育技术常见的研究范式。涉及实验组 vs 对照组对比、教学法干预或技术接入前后测。",
                "attack_vectors": [
                    {
                        "id": "qv_1",
                        "pitfall_title": "前测基准等价性与二语水平未严格控制 (Baseline Equivalence & Proficiency)",
                        "reviewer_attack": "What were the pre-study assessments on writing proficiency exactly? How do you ensure the groups were truly comparable before the intervention?",
                        "danger_level": "极高 (High Risk)",
                        "target_section": "Methods -> 3.1 Participants / Baseline",
                        "preemptive_strategy": "不要仅用“同专业同班级”带过。必须在投稿前明确汇报前测工具（如 TEM-4 标准化成绩或受控作文基线得分），并提供单因素方差分析或 t 检验的 p > .05 齐同性证据。",
                        "preemptive_template": "To ensure baseline equivalence between the experimental and control cohorts, pre-intervention writing proficiency was assessed using [Instrument, e.g., standardized TEM-4 scores]. An independent-samples t-test indicated no statistically significant difference between the two groups (t(df) = ..., p = .xxx, Cohen's d = .xx), confirming their baseline comparability.",
                        "history_evidence": "在 System (Rev #1, R1-C01) 和 CALL (Rev #1, Table R3) 中，审稿人均点名追问前测基线细节。作者通过增补标准化测验说明与非被试排除标准成功化解。"
                    },
                    {
                        "id": "qv_2",
                        "pitfall_title": "干预保真度与任务时间线模糊 (Treatment Fidelity & Timeline)",
                        "reviewer_attack": "The exact procedure and instructional conditions remain vague. How was feedback delivered, and did students in both groups spend equal task time?",
                        "danger_level": "中高 (Moderate-High)",
                        "target_section": "Methods -> 3.2 Pedagogical Intervention & Procedure",
                        "preemptive_strategy": "在正文中绘制一张清晰的干预周期时间轴（Timeline），精确注明每周课时、任务类型、同伴反馈交互轮次，并说明授课教师是否一致（控制教师效应）。",
                        "preemptive_template": "The intervention spanned [X] weeks, comprising [Y] sequential cycles. Both groups were taught by the same instructor following identical syllabi, with the sole independent variable being [e.g., AI-assisted vs. unassisted peer feedback]. Task time was equalized across conditions (90 mins/week).",
                        "history_evidence": "在 ILE (Rev #1) 和 System (Rev #1, R1-C02) 中，作者补充了教学环境详细说明与具体课时安排表，审稿人在 Rev 2 中完全认可。"
                    }
                ]
            },
            {
                "id": "small_sample",
                "name": "小样本 / 案例研究 / 探索性研究 (Small Sample / Case Study, N < 30)",
                "icon": "users",
                "tag": "质性与案例必测",
                "description": "针对特定班级或深度追踪案例的研究，样本量通常有限（N=4~30），极易招致统计效力质疑。",
                "attack_vectors": [
                    {
                        "id": "ss_1",
                        "pitfall_title": "样本量过小导致结论外推性与代表性受限 (Generalizability Constraint)",
                        "reviewer_attack": "Consider expanding the sample size and diversity of participants to enhance the generalizability of findings. Four participants seem too few to draw robust claims.",
                        "danger_level": "极高 (Fatal if unaddressed)",
                        "target_section": "Introduction -> Rationale & Discussion -> Limitations",
                        "preemptive_strategy": "投稿前在正文坦承案例研究的质性深度定位，强调研究目的在于机制探索（Explanatory Power）与情境化深描（Thick Description），而非统计学大样本外推。并在 Limitations 中诚恳阐明。",
                        "preemptive_template": "While the small cohort (N = [X]) precludes statistical generalization to broader populations, this study deliberately adopted an intensive multiple-case design to capture the nuanced, longitudinal trajectory of [...]. As Yin (2018) notes, the methodological merit of case inquiry lies in analytical rather than statistical generalization.",
                        "history_evidence": "在 HSSC (Reviewer 6) 与 Kybernetes (Rev #1, R1-C02) 中，面对审稿人对 N=4 和小样本的猛烈攻击，作者以“案例研究先天的理论定位”为盾，顺利过审。"
                    },
                    {
                        "id": "ss_2",
                        "pitfall_title": "被试流失与自愿性招募偏差未作说明 (Attrition & Volunteer Bias)",
                        "reviewer_attack": "Why did only a fraction of initial volunteers complete all stages? How does attrition affect the validity of data?",
                        "danger_level": "中等 (Medium)",
                        "target_section": "Methods -> Participants & Data Collection",
                        "preemptive_strategy": "提前在正文注明最初招募人数、完整数据筛选标准（Inclusion/Exclusion criteria），说明流失原因（如出勤、技术设备故障），避免审稿人误以为作者在刻意挑选数据。",
                        "preemptive_template": "Originally, [X] learners volunteered for the project. Following the data curation protocol (requiring completion of all writing drafts and reflection logs), [Y] participants constituted the final analytic dataset. Reasons for attrition were primarily logistical rather than intervention-induced.",
                        "history_evidence": "在 HSSC 和 CALL 的返修中均有此条。作者提前交代数据完整性筛除规则可直接避免此性质疑。"
                    }
                ]
            },
            {
                "id": "genai_tech",
                "name": "生成式 AI / 自动技术介入研究 (GenAI / LLM / Automated Feedback)",
                "icon": "bot",
                "tag": "前沿技术与伦理必测",
                "description": "涉及 ChatGPT, LLM, 自动写作评价（AWE/AWCF）等技术介入，当前 SSCI 顶刊审稿人对此关切极大。",
                "attack_vectors": [
                    {
                        "id": "gt_1",
                        "pitfall_title": "提示词黑箱与复现性缺陷 (Prompt Engineering Transparency)",
                        "reviewer_attack": "The prompting template used for AI feedback should be provided as an appendix. Without exact prompts, the study lacks methodological transparency.",
                        "danger_level": "极高 (Standard requirement in 2024+)",
                        "target_section": "Methods -> GenAI Implementation & Appendix",
                        "preemptive_strategy": "千万不要只在正文用自然语言描述'我们让 ChatGPT 提供反馈'！必须在初稿附录（Appendix）中完整附上 System Prompt, User Prompt 模版，并正文说明提示词如何对齐评价标准。",
                        "preemptive_template": "To ensure methodological transparency and ecological reproducibility, the verbatim prompt templates and regulatory guidelines employed for AI-mediated feedback are provided in Appendix 1. The prompt was specifically configured to guide learners on [dimensions] without directly rewriting.",
                        "history_evidence": "在 System (Rev #1, R1-C04) 中，Reviewer 1 明确要求将 Prompt 移至附录，作者在返修中增补后获得高度评价。"
                    },
                    {
                        "id": "gt_2",
                        "pitfall_title": "高校机构 AI 政策与学术诚信伦理模糊 (Institutional Policy & Ethics)",
                        "reviewer_attack": "The section on GenAI policy is vague, simply referring to TEQSA guidelines without specifying how student autonomy and ethics were safeguarded.",
                        "danger_level": "高 (High)",
                        "target_section": "Methods -> 3.X GenAI Policy & Ethical Safeguards",
                        "preemptive_strategy": "设立独立的 Institutional GenAI Policy 小节，引用 TEQSA (2023) 或院校政策，强调以学习者为中心的辅助性定位（Augmentation, not substitution），明确告知知情同意与数据隐私。",
                        "preemptive_template": "This study adhered to the TEQSA (2023) guidelines on generative AI in education. GenAI was purposefully introduced as a cognitive scaffold to augment, rather than replace, student-generated feedback. Institutional ethical approval and written informed consent were obtained prior to data collection.",
                        "history_evidence": "在 System (Rev #1, R1-C03) 和 EAIT 中，作者专门将 AI 政策从文献综述重构为研究方法的第一小节，彻底打消了审稿人的伦理顾虑。"
                    }
                ]
            },
            {
                "id": "measurement_survey",
                "name": "问卷调查与量表测量 (Questionnaire / Survey / Measurement Scales)",
                "icon": "clipboard-check",
                "tag": "构念效度必测",
                "description": "涉及李克特量表、自我感受（Perception）、元认知、投入度（Engagement）或认知负荷测量。",
                "attack_vectors": [
                    {
                        "id": "ms_1",
                        "pitfall_title": "量表信效度与操作化定义缺失 (Reliability, Validity & Operationalization)",
                        "reviewer_attack": "The questionnaire dimensions remain under-theorized. What are the Cronbach's alpha coefficients for each subscale? How was construct validity verified?",
                        "danger_level": "极高 (Fatal if neglected)",
                        "target_section": "Methods -> Instruments / Measurements",
                        "preemptive_strategy": "每个量表维度必须给出学术出处（Adapted from Author, Year），汇报样本实测内部一致性信度（Cronbach's $\alpha$），如有条件简要说明双语翻译与回译（Back-translation）流程。",
                        "preemptive_template": "The questionnaire was adapted from [Author, Year] and contextualized for [Context]. The instrument comprises [X] subscales: [A] (α = .xx), [B] (α = .xx), and [C] (α = .xx). The overall scale demonstrated satisfactory internal consistency (Cronbach's α = .xx). Construct validity was verified via confirmatory factor analysis.",
                        "history_evidence": "在 Fip, CAEH 与 InJAL 中，审稿人均要求补充具体信度数值与题目样例，作者补全后直接过审。"
                    },
                    {
                        "id": "ms_2",
                        "pitfall_title": "单一自评偏差与三角互证缺失 (Self-report Bias & Triangulation)",
                        "reviewer_attack": "Self-perceptions may diverge substantially from actual application abilities. How did you triangulate self-reported survey findings?",
                        "danger_level": "高 (High)",
                        "target_section": "Methods & Discussion",
                        "preemptive_strategy": "若使用自评量表，正文 Methods 必须声明其为主观认知测量，并在结果与讨论中补充行为数据（如修改痕迹、反思日志、课堂测试表现）进行三角互证（Methodological Triangulation）。",
                        "preemptive_template": "To mitigate self-report bias, survey responses were triangulated with objective textual metrics (e.g., revision uptake tracking) and stimulated-recall interview data, thereby achieving methodological triangulation (Denzin, 2012).",
                        "history_evidence": "在 IJARPED（数字素养自我感受及应用）和 InJAL 中，作者通过文本修改记录成功抵御了“只测了感受、没测实际能力”的质疑。"
                    }
                ]
            },
            {
                "id": "theoretical_framing",
                "name": "理论框架与核心构念 (Theoretical Framework & Core Constructs)",
                "icon": "book-open",
                "tag": "宏观论证必测",
                "description": "审稿人最爱扣上的大帽子：'理论薄弱'、'概念混淆'、'讨论缺乏机制阐释'。",
                "attack_vectors": [
                    {
                        "id": "tf_1",
                        "pitfall_title": "核心术语混淆与操作化定义模糊 (Terminological Slippage & Ambiguity)",
                        "reviewer_attack": "The construct of 'uptake' remains under-theorized. It is introduced without a clear operational definition, and its link with 'revision' is not adequately explained.",
                        "danger_level": "高 (High)",
                        "target_section": "Literature Review -> 2.1 Theoretical Framework",
                        "preemptive_strategy": "在文献综述末尾建立专门的'概念透镜与操作化界定'段落，明确界定容易混淆的概念区别（如 Uptake vs Revision, Curriculum vs Course, Feedback Seeking vs SRL）。",
                        "preemptive_template": "In the context of the present inquiry, we operationalize [Construct A] as [...], distinguishing it from [Construct B] which denotes [...]. Grounded in [Theorist, Year]'s socio-cognitive model, this distinction is critical because [...].",
                        "history_evidence": "在 InJAL (Rev #2, R2-C04) 和 Fip (Reviewer 1, Term Use) 中，作者通过重新梳理概念层级界定，赢得了审稿人的称赞。"
                    },
                    {
                        "id": "tf_2",
                        "pitfall_title": "讨论部分仅仅复述数据，缺乏批判性机制提炼 (Descriptive vs Critical Discussion)",
                        "reviewer_attack": "Discussion lacks critical argumentations focusing on the reasons for the findings based on the conceptual framework.",
                        "danger_level": "极高 (Common Reason for Major Revision)",
                        "target_section": "Discussion -> Underlying Mechanisms",
                        "preemptive_strategy": "绝不在讨论部分只写'实验组分数显著高于对照组'。必须以理论框架中的机制（如认知负荷、社会共享调节、对话式反馈循环）解释'为什么会这样'，并指出与经典文献的分歧点。",
                        "preemptive_template": "Rather than merely mirroring the empirical gains, these findings can be theorized through the lens of [Theoretical Framework]. Specifically, the synergy between [Factor A] and [Factor B] mitigated cognitive extraneous load, thereby facilitating [...]. This contrasts with [Prior Scholar]'s assertion that [...], suggesting that context plays a moderating role.",
                        "history_evidence": "在 CALL (Rev #1, Comment 2) 和 LLT (Rev #2, Comment 2-3) 中，作者通过系统性引入理论解释透镜重写讨论小节，顺利拿到 Accept。"
                    }
                ]
            },
            {
                "id": "qualitative_coding",
                "name": "质性编码与交互话语分析 (Qualitative Coding & Interaction Analysis)",
                "icon": "message-square-code",
                "tag": "质性严谨性必测",
                "description": "涉及学生同伴反馈评语分类、访谈转写、认知网络分析（ENA）或互动话语分析。",
                "attack_vectors": [
                    {
                        "id": "qc_1",
                        "pitfall_title": "质性编码者一致性信度未报告 (Inter-coder Reliability Missing)",
                        "reviewer_attack": "How was the coding scheme developed and validated? No inter-rater reliability statistic (e.g., Cohen's Kappa) is provided for the feedback categorization.",
                        "danger_level": "高 (Standard Reviewer Checklist)",
                        "target_section": "Methods -> Data Analysis / Coding Scheme",
                        "preemptive_strategy": "初稿中必须明确报告：独立编码者人数、抽样百分比（通常 20%~30%）、经过多轮讨论后的 Cohen's Kappa 系数（建议 > .80），并在正文中提供典型编码样例表。",
                        "preemptive_template": "To ensure analytic rigor, two raters independently coded a randomly selected 20% subset of the qualitative corpus ([N] segments). Inter-coder reliability was substantial (Cohen's Kappa = .86). Discrepancies were resolved through deliberative consensus with a third senior researcher.",
                        "history_evidence": "在 System, CALL, ILE 的返修中均有此要求。初稿提前写上，审稿人直接无刺可挑。"
                    }
                ]
            }
        ]
    }

def main():
    print("=" * 65)
    print("  AcadRebuttal · 自适应学术工作区智能扫描与数据湖编译")
    print(f"  根工作区: {WORKSPACE_DIR}")
    print(f"  输出文件: {OUTPUT_JSON}")
    print("=" * 65)

    discovered_projects = scan_workspace()
    print(f"\n[发现] 共动态识别到 {len(discovered_projects)} 个论文工作区目录:")
    for p in discovered_projects:
        print(f"  - [{p['journal']}] {p['topic']} ({len(p['files'])} 份返修信)")

    all_records = []
    seq_id = 1

    for p in discovered_projects:
        jname = p['journal']
        jtier = p['journal_tier']
        topic = p['topic']

        for r_name, full_path, rel_path in p['files']:
            if rel_path.endswith('.pdf'):
                raw_items = parse_pdf_adaptive(full_path)
            else:
                raw_items = parse_docx_adaptive(full_path)

            print(f"  [提取] {jname} ({r_name}): 成功抽取 {len(raw_items)} 条问答对 <- {os.path.basename(rel_path)}")

            for rev_role, comment_txt, resp_txt in raw_items:
                comment_clean = comment_txt.strip()
                resp_clean = resp_txt.strip()
                if len(comment_clean) < 10 or len(resp_clean) < 10:
                    continue

                category, strategy = classify_text(comment_clean, resp_clean)
                golden_move = extract_golden_move(resp_clean)
                keywords = extract_keywords(comment_clean + " " + resp_clean)

                rec = {
                    "id": f"REV_{seq_id:04d}",
                    "journal": jname,
                    "journal_tier": jtier,
                    "paper_topic": topic,
                    "revision_round": r_name,
                    "reviewer_role": rev_role,
                    "category": category,
                    "strategy": strategy,
                    "comment_text": comment_clean,
                    "response_text": resp_clean,
                    "golden_move": golden_move,
                    "keywords": keywords,
                    "source_path": rel_path
                }
                all_records.append(rec)
                seq_id += 1

    print(f"\n[OK] 全自动扫描与抽取完成，全库共计: {len(all_records)} 条学术对局")

    # Rebuild Knowledge Graph
    kg = build_knowledge_graph(all_records)
    print(f"[OK] 动态编织知识拓扑网络: {len(kg['nodes'])} 个节点, {len(kg['links'])} 条连线")

    # Compute Multi-dimensional Analytics
    journals_dist = {}
    categories_dist = {}
    strategies_dist = {}
    rounds_dist = {}
    for r in all_records:
        journals_dist[r['journal']] = journals_dist.get(r['journal'], 0) + 1
        categories_dist[r['category']] = categories_dist.get(r['category'], 0) + 1
        strategies_dist[r['strategy']] = strategies_dist.get(r['strategy'], 0) + 1
        rounds_dist[r['revision_round']] = rounds_dist.get(r['revision_round'], 0) + 1

    analytics = {
        "total_records": len(all_records),
        "journal_distribution": journals_dist,
        "category_distribution": categories_dist,
        "strategy_distribution": strategies_dist,
        "round_distribution": rounds_dist
    }

    bundle = {
        "version": "2.1.0",
        "description": "AcadRebuttal Dynamic Academic Data Lake",
        "metadata": {
            "total_records": len(all_records),
            "total_journals": len(journals_dist),
            "categories": sorted(list(categories_dist.keys())),
            "strategies": sorted(list(strategies_dist.keys())),
            "journals": sorted(list(journals_dist.keys())),
            "rounds": sorted(list(rounds_dist.keys()))
        },
        "records": all_records,
        "knowledge_graph": kg,
        "analytics": analytics,
        "strategy_vault": get_strategy_vault(all_records),
        "scanner_data": get_presubmission_scanner_data()
    }

    os.makedirs(os.path.dirname(OUTPUT_JSON), exist_ok=True)
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(bundle, f, ensure_ascii=False, indent=2)

    print(f"[OK] 已将最新语料全量编译输出至前端数据湖: {OUTPUT_JSON}")

if __name__ == '__main__':
    main()
